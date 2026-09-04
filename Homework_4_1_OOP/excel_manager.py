"""Functions for reading store data from Excel files."""

from openpyxl import load_workbook

from customer import Customer
from product import Product


def read_data_from_excel(file_name: str) -> tuple[list[Product], list[Customer]]:
    """Read products and customers from an Excel file."""
    workbook = load_workbook(file_name)
    sheet = workbook.active

    products: list[Product] = []
    customers: list[Customer] = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        object_type = row[0]

        if object_type == "PRODUCT":
            name = row[1]
            category = row[2]
            price = float(row[3])
            quantity = int(row[4])

            product = Product(name, category, price, quantity)

            products.append(product)

        elif object_type == "CUSTOMER":
            name = row[1]
            email = row[2]

            customer = Customer(name, email)
            customers.append(customer)

    return products, customers
